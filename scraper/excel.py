"""XLSX export (openpyxl): Data / Summary / Opportunities / Run Info sheets."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from scraper import TOOL_NAME
from scraper.analysis import CategorySummary, Opportunity
from scraper.models import Book

GBP = '"£"#,##0.00'
PCT = "0.0%"
TABLE_STYLE = "TableStyleMedium2"
TITLE_FONT = Font(size=14, bold=True, color="1F3864")
NOTE_FONT = Font(size=9, italic=True, color="595959")
DEMO_FILL = PatternFill("solid", fgColor="FCE4D6")
LIMIT_FONT = Font(size=9, italic=True, color="C55A11")

# A cell whose text starts with one of these is read as a formula by Excel/LibreOffice when a CSV is
# opened ("CSV injection"). Scraped text is untrusted, so it must always stay plain text.
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def csv_safe(value: Any) -> Any:
    """CSV has no cell types: prefix risky text with an apostrophe so spreadsheets show it as text."""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def _add_table(ws: Worksheet, name: str, header_row: int, n_rows: int, n_cols: int) -> None:
    """Wrap a block in an Excel Table: banded rows + filter buttons on every header."""
    ref = f"A{header_row}:{get_column_letter(n_cols)}{header_row + max(n_rows, 1)}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name=TABLE_STYLE, showRowStripes=True)
    ws.add_table(table)


def _write_block(ws: Worksheet, start_row: int, headers: Sequence[str],
                 rows: Iterable[Sequence[Any]], widths: Sequence[float],
                 formats: dict[int, str] | None = None, link_col: int | None = None) -> int:
    """Write headers + rows starting at `start_row`. Returns number of data rows written."""
    formats = formats or {}
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        ws.cell(row=start_row, column=c, value=h)
        ws.column_dimensions[get_column_letter(c)].width = w
    n = 0
    for n, row in enumerate(rows, start=1):
        r = start_row + n
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if cell.data_type == "f":
                # openpyxl turns any string starting with "=" into a live formula. Scraped text never
                # is one: store it as a string cell, shown exactly as scraped.
                cell.data_type = "s"
            if c in formats:
                cell.number_format = formats[c]
            if link_col == c and value:
                cell.hyperlink = value
                cell.style = "Hyperlink"
    return n


def _sheet_header(ws: Worksheet, title: str, note: str) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = note
    ws["A2"].font = NOTE_FONT
    ws.row_dimensions[1].height = 22


def write_xlsx(path: Path, books: list[Book], summary: list[CategorySummary],
               opportunities: list[Opportunity], run_info: dict[str, Any]) -> Path:
    wb = Workbook()
    wb.properties.creator = TOOL_NAME
    wb.properties.title = "Competitor catalog snapshot (DEMO)"

    # ---------------------------------------------------------------- Data
    ws = wb.active
    ws.title = "Data"
    headers = ["Title", "Category", "Price (GBP)", "Rating", "Availability", "In stock", "Product URL",
               "Scraped at (UTC)"]
    rows = []
    for b in books:
        r = b.as_row()
        rows.append([r["title"], r["category"], r["price_gbp"], r["rating"], r["availability"],
                     "Yes" if r["in_stock"] else "No", r["url"], r["scraped_at"]])
    n = _write_block(ws, 1, headers, rows, widths=[58, 20, 13.5, 9, 13, 10.5, 62, 19],
                     formats={3: GBP, 4: "0"}, link_col=7)
    _add_table(ws, "tblBooks", 1, n, len(headers))
    ws.freeze_panes = "A2"

    # ------------------------------------------------------------- Summary
    ws = wb.create_sheet("Summary")
    _sheet_header(ws, "Catalog summary by category",
                  f"DEMO data · source: {run_info['source']} · collected {run_info['collected_at']}")
    if run_info.get("limit_note"):  # --max-pages: the Titles column only counts the pages read
        ws["A3"] = f"Page limit: {run_info['limit_note']}"
        ws["A3"].font = LIMIT_FONT
    headers = ["Category", "Titles", "Avg price", "Median price", "Min price", "Max price", "Avg rating",
               "In stock %", "Opportunities"]
    rows = [[s.category, s.titles, s.avg_price, s.median_price, s.min_price, s.max_price, s.avg_rating,
             s.in_stock_pct, s.opportunities] for s in summary]
    start = 4
    n = _write_block(ws, start, headers, rows, widths=[24, 10, 12, 15, 12, 12, 12.5, 12.5, 15.5],
                     formats={3: GBP, 4: GBP, 5: GBP, 6: GBP, 7: "0.00", 8: PCT})
    _add_table(ws, "tblSummary", start, n, len(headers))
    if n:
        # Data bar makes the catalog mix readable at a glance.
        ws.conditional_formatting.add(
            f"B{start + 1}:B{start + n}",
            DataBarRule(start_type="num", start_value=0, end_type="max", color="5B9BD5"))
        # Totals row sits right under the table (outside it, so filters/sorts don't move it).
        total_r = start + n + 1
        total_titles = sum(s.titles for s in summary)
        all_prices = [b.price_gbp for b in books]
        values = ["All categories", total_titles,
                  round(sum(all_prices) / len(all_prices), 2) if all_prices else None,
                  run_info.get("median_price"), min(all_prices, default=None), max(all_prices, default=None),
                  run_info.get("avg_rating"), run_info.get("in_stock_pct"), len(opportunities)]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=total_r, column=c, value=v)
            cell.font = Font(bold=True)
            cell.number_format = {3: GBP, 4: GBP, 5: GBP, 6: GBP, 7: "0.00", 8: PCT}.get(c, "General")
    ws.freeze_panes = f"A{start + 1}"

    # ------------------------------------------------------- Opportunities
    ws = wb.create_sheet("Opportunities")
    _sheet_header(ws, "Opportunities: rating >= 4 and price below the category median",
                  "Well-rated titles that are cheaper than their category peers. "
                  "Sorted by rating, then by % below median.")
    headers = ["Title", "Category", "Price (GBP)", "Category median", "Below median", "Rating", "Product URL"]
    rows = [[o.book.title, o.book.category, o.book.price_gbp, o.category_median, round(o.below_median_pct, 4),
             o.book.rating, str(o.book.url)] for o in opportunities]
    n = _write_block(ws, start, headers, rows, widths=[58, 20, 13.5, 18, 15, 9, 62],
                     formats={3: GBP, 4: GBP, 5: PCT, 6: "0"}, link_col=7)
    _add_table(ws, "tblOpportunities", start, n, len(headers))
    ws.freeze_panes = f"A{start + 1}"

    # ------------------------------------------------------------ Run Info
    ws = wb.create_sheet("Run Info")
    _sheet_header(ws, "Run information", "Generated automatically by the scraper. Numbers come from run_log.json.")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    for i, (k, v) in enumerate(run_info.get("table", []), start=4):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value=v)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    demo_row = 4 + len(run_info.get("table", [])) + 1
    ws.cell(row=demo_row, column=1, value="DEMO").font = Font(bold=True, color="C55A11")
    note = ws.cell(row=demo_row, column=2, value=(
        "Sample output from a public scraping sandbox. Prices and ratings on books.toscrape.com are "
        "randomly assigned and have no real meaning."))
    note.fill = DEMO_FILL
    note.alignment = Alignment(wrap_text=True)

    # Open on Summary: its first lines say DEMO and where the data comes from. The Data sheet stays
    # first in the tab order; only one tab may be selected, or Excel opens the two sheets grouped.
    wb.active = wb.sheetnames.index("Summary")
    for sheet in wb.worksheets:
        sheet.sheet_view.tabSelected = sheet.title == "Summary"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
