"""Render a generated .xlsx into a static, spreadsheet-like HTML page (used for the README media).

Everything shown comes from the file itself via openpyxl: values, number formats, column widths,
fonts, fills, hyperlinks, frozen rows, table styles (banding/header of TableStyleMedium2 is
emulated because Excel applies it at render time) and data-bar conditional formats.

    python scripts/xlsx_preview.py examples/output/books.xlsx -o tmp/books_preview.html
"""

from __future__ import annotations

import argparse
import html
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

MAX_ROWS = 1200
DEFAULT_COL_W = 8.43
# TableStyleMedium2 (Office theme): header = accent1, banded rows = accent1 tint 80%
TABLE_HEADER_BG, TABLE_HEADER_FG, TABLE_STRIPE = "#4472C4", "#FFFFFF", "#D9E1F2"


def col_px(width: float | None) -> int:
    # Excel column width (in "0" characters of Calibri 11) -> pixels
    w = width or DEFAULT_COL_W
    return int(w * 7 + 5)


def fmt_value(cell: Cell) -> str:
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    nf = cell.number_format or "General"
    if isinstance(v, (int, float)):
        if "%" in nf:
            decimals = len(nf.split(".")[1].rstrip("%")) if "." in nf else 0
            return f"{v * 100:.{decimals}f}%"
        m = re.search(r"0\.(0+)", nf)
        decimals = len(m.group(1)) if m else 0
        prefix = "£" if "£" in nf else ""
        if nf == "General":
            return f"{v:g}" if isinstance(v, float) else str(v)
        return f"{prefix}{v:,.{decimals}f}" if "," in nf else f"{prefix}{v:.{decimals}f}"
    return str(v)


def color_of(c) -> str | None:
    """openpyxl Color -> '#RRGGBB' (theme/indexed colors are ignored: we only use explicit RGB)."""
    if c is None or c.type != "rgb" or not isinstance(c.rgb, str):
        return None
    rgb = c.rgb[-6:]
    return None if c.rgb in ("00000000",) else f"#{rgb}"


def table_regions(ws: Worksheet) -> list[tuple[int, int, int, int]]:
    out = []
    for t in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(t.ref)
        out.append((min_col, min_row, max_col, max_row))
    return out


def databars(ws: Worksheet) -> dict[tuple[int, int], tuple[float, str]]:
    """(row, col) -> (fraction 0..1, color) for every cell covered by a DataBar rule."""
    bars: dict[tuple[int, int], tuple[float, str]] = {}
    for rng in ws.conditional_formatting:
        for rule in rng.rules:
            if rule.type != "dataBar" or rule.dataBar is None:
                continue
            color = f"#{rule.dataBar.color.rgb[-6:]}"
            for cr in rng.sqref.ranges:
                min_col, min_row, max_col, max_row = range_boundaries(cr.coord)
                cells = [(r, c) for r in range(min_row, max_row + 1) for c in range(min_col, max_col + 1)]
                vals = [ws.cell(r, c).value for r, c in cells]
                top = max((v for v in vals if isinstance(v, (int, float))), default=0) or 1
                for (r, c), v in zip(cells, vals):
                    if isinstance(v, (int, float)):
                        bars[(r, c)] = (max(0.0, v / top), color)
    return bars


def render_sheet(ws: Worksheet, idx: int) -> str:
    max_row = min(ws.max_row, MAX_ROWS)
    max_col = ws.max_column
    regions = table_regions(ws)
    bars = databars(ws)
    frozen_rows = (ws.freeze_panes and int(re.sub(r"\D", "", ws.freeze_panes)) - 1) or 0

    widths = [col_px(ws.column_dimensions[get_column_letter(c)].width) for c in range(1, max_col + 1)]
    cols = "".join(f'<col style="width:{w}px">' for w in [44] + widths)
    head = "".join(f"<th>{get_column_letter(c)}</th>" for c in range(1, max_col + 1))
    body = []
    top = 22  # sticky offset below the column-letter header
    for r in range(1, max_row + 1):
        h_pt = ws.row_dimensions[r].height
        h = int(h_pt * 4 / 3) if h_pt else 20
        sticky = f' class="frozen" style="top:{top}px"' if r <= frozen_rows else ""
        if r <= frozen_rows:
            top += h
        tds = [f'<td class="rh">{r}</td>']
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            styles, classes = [], []
            text = html.escape(fmt_value(cell))
            region = next((g for g in regions if g[0] <= c <= g[2] and g[1] <= r <= g[3]), None)
            if region:
                if r == region[1]:
                    styles += [f"background:{TABLE_HEADER_BG}", f"color:{TABLE_HEADER_FG}", "font-weight:700"]
                    classes.append("thd")
                    text += '<span class="flt">&#9662;</span>'
                elif (r - region[1]) % 2 == 1:
                    styles.append(f"background:{TABLE_STRIPE}")
            fill = cell.fill
            if fill is not None and fill.fill_type == "solid" and color_of(fill.fgColor):
                styles.append(f"background:{color_of(fill.fgColor)}")
            f = cell.font
            if f is not None:
                if f.b:
                    styles.append("font-weight:700")
                if f.i:
                    styles.append("font-style:italic")
                if f.sz and float(f.sz) != 11:
                    styles.append(f"font-size:{float(f.sz) * 4 / 3:.1f}px")
                if color_of(f.color) and not (region and r == region[1]):
                    styles.append(f"color:{color_of(f.color)}")
            if cell.hyperlink:
                classes.append("link")
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                classes.append("num")
            if cell.alignment is not None and cell.alignment.wrap_text:
                classes.append("wrap")
            elif isinstance(cell.value, str) and not region and c < max_col and ws.cell(r, c + 1).value is None:
                classes.append("spill")  # like Excel: text runs over empty neighbours instead of clipping
            if (r, c) in bars:
                frac, color = bars[(r, c)]
                styles.append(f"background-image:linear-gradient(90deg,{color} 0%,#fff {frac * 100:.1f}%,"
                              f"transparent {frac * 100:.1f}%)")
            cls = f' class="{" ".join(classes)}"' if classes else ""
            sty = f' style="{";".join(styles)}"' if styles else ""
            tds.append(f"<td{cls}{sty}>{text}</td>")
        body.append(f'<tr{sticky} style="height:{h}px">{"".join(tds)}</tr>')
    hidden = "" if idx == 0 else " hidden"
    # table-layout:fixed only honours <col> widths when the table itself has an explicit width
    total_w = 44 + sum(widths)
    return (f'<div class="sheet{hidden}" data-sheet="{html.escape(ws.title)}">'
            f'<table class="grid" style="width:{total_w}px">'
            f'<colgroup>{cols}</colgroup><thead><tr><th class="corner"></th>{head}</tr></thead>'
            f'<tbody>{"".join(body)}</tbody></table></div>')


PAGE = """<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{name} preview</title>
<style>
  :root {{ --line:#d4d4d4; --hdr:#f3f3f3; --ink:#1f1f1f; }}
  * {{ box-sizing:border-box; }}
  html,body {{ margin:0; height:100%; background:#fff; color:var(--ink);
    font:14.67px Calibri, Carlito, "Segoe UI", Arial, sans-serif; }}
  .app {{ display:flex; flex-direction:column; height:100vh; }}
  .titlebar {{ background:#2b2b2b; color:#fff; padding:9px 14px; display:flex; align-items:center; gap:12px;
    font:13px "Segoe UI", Arial, sans-serif; }}
  .titlebar b {{ font-size:14px; }}
  .badge {{ background:#eb6834; border-radius:10px; padding:2px 9px; font-weight:700; font-size:11px; }}
  .titlebar .src {{ margin-left:auto; color:#bdbdbd; }}
  .fbar {{ display:flex; border-bottom:1px solid var(--line); font:13px "Segoe UI", Arial, sans-serif; }}
  .fbar .nm {{ width:90px; padding:5px 8px; border-right:1px solid var(--line); }}
  .fbar .fx {{ padding:5px 8px; color:#666; border-right:1px solid var(--line); font-style:italic; }}
  .fbar .val {{ padding:5px 8px; white-space:nowrap; overflow:hidden; }}
  .wrapgrid {{ flex:1; overflow:auto; position:relative; }}
  .sheet.hidden {{ display:none; }}
  table.grid {{ border-collapse:separate; border-spacing:0; table-layout:fixed; }}
  .grid td, .grid th {{ border-right:1px solid var(--line); border-bottom:1px solid var(--line); padding:0 5px;
    white-space:nowrap; overflow:hidden; text-overflow:clip; height:20px; }}
  .grid thead th {{ position:sticky; top:0; z-index:3; background:var(--hdr); color:#444; font:12px "Segoe UI",
    Arial, sans-serif; text-align:center; height:22px; }}
  .grid td.rh {{ position:sticky; left:0; z-index:2; background:var(--hdr); color:#444; text-align:center;
    font:12px "Segoe UI", Arial, sans-serif; }}
  .grid th.corner {{ left:0; z-index:4; }}
  .grid tr.frozen td {{ position:sticky; z-index:1; background-color:#fff; }}
  .grid tr.frozen td.rh {{ z-index:3; background:var(--hdr); }}
  .grid tr.frozen:last-of-type td, .grid tr.frozen + tr:not(.frozen) td {{ }}
  td.num {{ text-align:right; }}
  td.link {{ color:#0563C1; text-decoration:underline; }}
  td.spill {{ overflow:visible; position:relative; z-index:2; }}
  tr.frozen td.spill {{ z-index:2; }}
  td.wrap {{ white-space:normal; line-height:1.25; padding-top:2px; padding-bottom:2px; }}
  td.thd {{ position:relative; padding-right:22px; }}
  .flt {{ position:absolute; right:3px; top:50%; transform:translateY(-50%); width:15px; height:15px; line-height:13px; text-align:center;
    font-size:10px; color:#444; background:#fff; border:1px solid #b9b9b9; border-radius:2px; }}
  .tabs {{ display:flex; align-items:flex-end; gap:2px; background:#f3f3f3; border-top:1px solid var(--line);
    padding:0 10px; height:34px; font:13px "Segoe UI", Arial, sans-serif; }}
  .tab {{ padding:7px 16px; cursor:pointer; color:#444; border-bottom:3px solid transparent; }}
  .tab.active {{ background:#fff; color:#1e7145; font-weight:600; border-bottom-color:#1e7145; }}
  .status {{ background:#f3f3f3; border-top:1px solid var(--line); padding:4px 12px; font:12px "Segoe UI", Arial,
    sans-serif; color:#555; display:flex; gap:24px; }}
</style></head><body><div class="app">
<div class="titlebar"><span class="badge">DEMO</span><b>{name}</b><span>preview rendered from the generated file</span>
<span class="src">{sheets_count} sheets · {rows_total} data rows</span></div>
<div class="fbar"><div class="nm" id="nm">A1</div><div class="fx">fx</div><div class="val" id="val"></div></div>
<div class="wrapgrid" id="grid">{sheets}</div>
<div class="tabs">{tabs}</div>
<div class="status"><span>Ready</span><span id="st"></span></div>
</div>
<script>
  const tabs = document.querySelectorAll('.tab');
  function show(name) {{
    document.querySelectorAll('.sheet').forEach(s => s.classList.toggle('hidden', s.dataset.sheet !== name));
    tabs.forEach(t => t.classList.toggle('active', t.dataset.sheet === name));
    const first = document.querySelector('.sheet:not(.hidden) tbody tr td:nth-child(2)');
    document.getElementById('val').textContent = first && first.firstChild ? first.firstChild.textContent : '';
    document.getElementById('grid').scrollTop = 0;
    const rows = document.querySelectorAll('.sheet:not(.hidden) tbody tr').length;
    document.getElementById('st').textContent = name + ' · ' + rows + ' rows';
  }}
  tabs.forEach(t => t.addEventListener('click', () => show(t.dataset.sheet)));
  window.showSheet = show;
  show(tabs[0].dataset.sheet);
</script></body></html>"""


def render(xlsx: Path) -> str:
    wb = load_workbook(xlsx)
    sheets = "".join(render_sheet(ws, i) for i, ws in enumerate(wb.worksheets))
    tabs = "".join(f'<div class="tab" data-sheet="{html.escape(ws.title)}">{html.escape(ws.title)}</div>'
                   for ws in wb.worksheets)
    rows_total = max(wb.worksheets[0].max_row - 1, 0)
    return PAGE.format(name=html.escape(xlsx.name), sheets=sheets, tabs=tabs, sheets_count=len(wb.worksheets),
                       rows_total=f"{rows_total:,}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("-o", "--out", type=Path, required=True)
    a = ap.parse_args()
    a.out.parent.mkdir(parents=True, exist_ok=True)
    a.out.write_text(render(a.xlsx), encoding="utf-8")
    print(a.out)


if __name__ == "__main__":
    main()
