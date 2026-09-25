"""Exporters end to end (CSV, XLSX, PDF, run_log.json) from a synthetic crawl result.

No browser and no network: build_outputs() only needs the in-memory outcome of a crawl.
"""

import csv
import json
import re
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from scraper.crawler import CrawlResult, CrawlSettings, PageEvent
from scraper.log import setup_logging
from scraper.models import Book, Category
from scraper.pipeline import CSV_COLUMNS, RunSettings, ScrapeOutcome, build_outputs
from scraper.robots import RobotsPolicy

NOW = datetime(2026, 1, 1, 12, 0, tzinfo=timezone.utc)
CATS = [Category(name=n, url=f"https://books.toscrape.com/catalogue/category/books/{n.lower()}_1/index.html")
        for n in ("Travel", "Poetry", "Mystery")]


def _book(title: str, category: str, price: float, rating: int, stock: bool = True) -> Book:
    return Book(title=title, category=category, price_gbp=price, rating=rating,
                availability="In stock" if stock else "Out of stock", in_stock=stock,
                url=f"https://books.toscrape.com/catalogue/{title.lower().replace(' ', '-')}_1/index.html",
                listing_page=1, scraped_at=NOW)


def _outcome(books: list[Book], failed_page: bool = False) -> ScrapeOutcome:
    pages = [PageEvent(category=c.name, page_no=1, url=c.url, status="ok", items=2, valid=2, attempts=1)
             for c in CATS[:2]]
    if failed_page:
        pages.append(PageEvent(category="Mystery", page_no=1, url=CATS[2].url, status="failed", attempts=3,
                               error="TimeoutError: simulated"))
    result = CrawlResult(books=books, pages=pages, retries=2 if failed_page else 0)
    robots = RobotsPolicy("https://books.toscrape.com/robots.txt", 404, "no robots.txt (HTTP 404): test")
    return ScrapeOutcome(robots=robots, available=CATS, selected=CATS, result=result, started_at=NOW,
                         crawl_seconds=12.3)


BOOKS = [
    _book("Cheap Gem", "Travel", 12.5, 5),
    _book("Mid Travel", "Travel", 30.0, 3),
    _book("Pricey Travel", "Travel", 55.0, 4),
    _book("Poem Deal", "Poetry", 10.0, 4, stock=False),
    _book("Poem Luxe", "Poetry", 48.0, 2),
]


@pytest.fixture()
def run(tmp_path):
    """Run the exporters once; a duplicate row and a failed page exercise the 'partial' path."""
    settings = RunSettings(out_dir=tmp_path, crawl=CrawlSettings())
    log = setup_logging(None, "WARNING")
    run_log = build_outputs(_outcome(BOOKS + [BOOKS[0]], failed_page=True), settings, log, wall_t0=0.0)
    return tmp_path, run_log


def test_csv_has_bom_header_and_deduplicated_rows(run):
    out, _ = run
    raw = (out / "books.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # utf-8-sig: Excel opens £ and accents correctly
    rows = list(csv.DictReader((out / "books.csv").open(encoding="utf-8-sig")))
    assert list(rows[0].keys()) == CSV_COLUMNS
    assert len(rows) == len(BOOKS)  # the duplicate URL was dropped
    assert rows[0]["price_gbp"] == "12.5" and rows[0]["rating"] == "5"


def test_xlsx_sheets_tables_and_formats(run):
    out, _ = run
    wb = load_workbook(out / "books.xlsx")
    assert wb.sheetnames == ["Data", "Summary", "Opportunities", "Run Info"]

    data = wb["Data"]
    assert "tblBooks" in data.tables and data.tables["tblBooks"].ref == f"A1:H{len(BOOKS) + 1}"
    assert data.freeze_panes == "A2"
    assert data["C2"].value == 12.5 and "£" in data["C2"].number_format
    assert data["G2"].hyperlink is not None  # product URL is clickable

    summary = wb["Summary"]
    # Header on row 4, one row per category, sorted by title count (Travel 3, Poetry 2)
    assert [summary.cell(row=r, column=1).value for r in (5, 6)] == ["Travel", "Poetry"]
    assert summary["B5"].value == 3 and summary["C5"].value == round((12.5 + 30 + 55) / 3, 2)
    assert summary["E5"].value == 12.5 and summary["F5"].value == 55.0

    # Opportunity rule: rating >= 4 and price below the category median
    # Travel median 30 -> Cheap Gem (5*, 12.50); Pricey Travel (4*) is above it. Poetry median 29 -> Poem Deal.
    opp = wb["Opportunities"]
    titles = [opp.cell(row=r, column=1).value for r in range(5, opp.max_row + 1)]
    assert titles == ["Cheap Gem", "Poem Deal"]


def test_pdf_is_valid_and_has_three_pages(run):
    out, _ = run
    pdf = (out / "report.pdf").read_bytes()
    assert pdf.startswith(b"%PDF-") and pdf.rstrip().endswith(b"%%EOF")
    # reportlab writes one "/Type /Page" dictionary per page (the "/Pages" tree node is excluded)
    assert len(re.findall(rb"/Type /Page(?!s)", pdf)) == 3
    assert (out / "charts" / "avg_price_by_category.png").stat().st_size > 5_000
    assert (out / "charts" / "rating_distribution.png").stat().st_size > 5_000


def test_run_log_reports_partial_run(run):
    out, returned = run
    saved = json.loads((out / "run_log.json").read_text(encoding="utf-8"))
    assert saved == returned
    assert saved["status"] == "partial"  # one page failed
    assert saved["pages"] == {"ok": 2, "failed": 1, "retries": 2, "skipped_by_robots": 0}
    assert saved["records"]["duplicates_removed"] == 1 and saved["records"]["exported"] == len(BOOKS)
    assert saved["errors"][0]["error"].startswith("TimeoutError")
    assert saved["kpis"]["opportunities"] == 2
    assert {"csv", "xlsx", "pdf", "run_log"} <= saved["outputs"].keys()


def test_no_valid_records_skips_xlsx_and_pdf(tmp_path):
    settings = RunSettings(out_dir=tmp_path, crawl=CrawlSettings())
    run_log = build_outputs(_outcome([]), settings, setup_logging(None, "WARNING"), wall_t0=0.0)
    assert run_log["status"] == "failed"
    assert (tmp_path / "books.csv").exists() and (tmp_path / "run_log.json").exists()
    assert not (tmp_path / "books.xlsx").exists() and not (tmp_path / "report.pdf").exists()
